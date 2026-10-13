"""Persist explicitly accepted, quality-gate-approved test-design artifacts."""

import hashlib
import json
import os
import re
from datetime import UTC, datetime
from pathlib import Path

from app.agents.context_converter_agent import ContextConverterAgent
from app.agents.test_case_validator import TestCaseValidatorAgent, ValidationReport
from app.models import (
    AcceptanceReceipt,
    AcceptedArtifact,
    AcceptSuiteRequest,
    ExecutionMode,
    ExportFormat,
    GenerateRequest,
    ManualArtifactFormat,
    TestSuite,
)

_SAFE_NAME = re.compile(r"[^a-z0-9]+")


class AcceptanceError(ValueError):
    """Raised when a suite cannot cross the explicit acceptance boundary."""


class AcceptedOutputService:
    def __init__(self, output_directory: Path) -> None:
        self.output_directory = output_directory
        self.converter = ContextConverterAgent()

    def accept(self, request: AcceptSuiteRequest) -> AcceptanceReceipt:
        if not request.validation.passed:
            raise AcceptanceError("Only a passing quality-gate suite can be accepted.")
        selected_ids = set(request.selected_case_ids)
        selected = [case for case in request.suite.test_cases if case.id in selected_ids]
        missing = selected_ids - {case.id for case in selected}
        if missing:
            raise AcceptanceError(f"Unknown selected test case IDs: {', '.join(sorted(missing))}")
        if not selected:
            raise AcceptanceError("Select at least one test case for acceptance.")
        context = self.validate_context(request.workflow_context)

        selected_suite = request.suite.model_copy(update={"test_cases": selected})
        structural_request = GenerateRequest(
            description=f"Accepted test suite for {request.suite.feature_name}.",
            output_format=request.suite.output_format,
        )
        if not TestCaseValidatorAgent().validate(structural_request, selected_suite).passed:
            raise AcceptanceError(
                "Selected cases no longer pass server-side structural validation."
            )

        self.output_directory.mkdir(parents=True, exist_ok=True)
        self._restrict_permissions(self.output_directory, 0o700)
        timestamp = datetime.now(UTC)
        stem = self._stem(request.suite.feature_name, timestamp)
        validation = ValidationReport.model_validate(request.validation.model_dump())
        artifacts: list[AcceptedArtifact] = []

        manual_cases = [case for case in selected if case.execution_mode == ExecutionMode.MANUAL]
        if manual_cases:
            manual_suite = request.suite.model_copy(update={"test_cases": manual_cases})
            export_format = (
                ExportFormat.CSV
                if request.manual_format == ManualArtifactFormat.CSV
                else ExportFormat.EXCEL
            )
            converted = self.converter.convert(manual_suite, validation, export_format)
            suffix = "csv" if export_format == ExportFormat.CSV else "xlsx"
            artifacts.append(
                self._store(
                    f"{stem}-manual-tests.{suffix}",
                    suffix,
                    len(manual_cases),
                    converted.content,
                )
            )

        automation_cases = [
            case for case in selected if case.execution_mode == ExecutionMode.AUTOMATION
        ]
        if automation_cases:
            automation_suite = request.suite.model_copy(update={"test_cases": automation_cases})
            converted = self.converter.convert(automation_suite, validation, ExportFormat.FEATURE)
            artifacts.append(
                self._store(
                    f"{stem}-automation-tests.feature",
                    "feature",
                    len(automation_cases),
                    converted.content,
                )
            )

        suite_hash = hashlib.sha256(request.suite.model_dump_json().encode("utf-8")).hexdigest()
        receipt = AcceptanceReceipt(
            accepted_at=timestamp.isoformat(),
            accepted_by=request.accepted_by.strip(),
            suite_hash=suite_hash,
            selected_case_ids=[case.id for case in selected],
            output_directory=str(self.output_directory),
            artifacts=artifacts,
        )
        receipt_content = json.dumps(receipt.model_dump(mode="json"), indent=2).encode("utf-8")
        self._store(
            f"{stem}-suite.json",
            "suite",
            len(selected),
            selected_suite.model_dump_json(indent=2).encode(),
        )
        if context:
            self._store(
                f"{stem}-workflow.json",
                "workflow",
                len(selected),
                json.dumps(context, indent=2).encode(),
            )
        self._store(f"{stem}-acceptance.json", "acceptance", len(selected), receipt_content)
        return receipt

    @staticmethod
    def validate_context(context: dict[str, object] | None) -> dict:
        """Preserve grounded stage evidence without promoting it into a business baseline."""
        if not context:
            return {}
        from app.workflow_models import ScenarioHandoff, StoryHandoff

        try:
            handoff = ScenarioHandoff if context.get("scenarios") else StoryHandoff
            parsed = handoff.model_validate(context)
            result = parsed.model_dump(mode="json")
            for key, ids in (
                ("story_approvals", {s.id for s in parsed.stories.stories}),
                (
                    "scenario_approvals",
                    {s.id for s in parsed.scenarios.scenarios}
                    if isinstance(parsed, ScenarioHandoff)
                    else set(),
                ),
            ):
                approvals = context.get(key, {})
                if not isinstance(approvals, dict) or not set(approvals) <= ids:
                    raise ValueError("Unknown stage approval IDs")
                for reviewer in approvals.values():
                    if not isinstance(reviewer, str) or not 2 <= len(reviewer.strip()) <= 100:
                        raise ValueError("Stage approvals require reviewer names")
                result[key] = approvals
            if context.get("ba_report"):
                from app.agents.business_analyst import BAReport

                result["ba_report"] = BAReport.model_validate(context["ba_report"]).model_dump(
                    mode="json"
                )
            if context.get("requirements_report"):
                from app.agents.requirements_validation import RequirementsReport

                result["requirements_report"] = RequirementsReport.model_validate(
                    context["requirements_report"]
                ).model_dump(mode="json")
            return result
        except ValueError as error:
            raise AcceptanceError("Workflow approval context is invalid: " + str(error)) from error

    def inventory(self) -> list[dict]:
        """Accepted files are approval history even when AI reuse is disabled."""
        result = []
        for path in sorted(self.output_directory.glob("*-acceptance.json"), reverse=True)[:500]:
            try:
                if path.is_symlink() or path.stat().st_size > 1_000_000:
                    continue
                receipt = AcceptanceReceipt.model_validate_json(path.read_bytes())
                stem = path.name.removesuffix("-acceptance.json")
                suite_path = self.output_directory / f"{stem}-suite.json"
                suite = (
                    TestSuite.model_validate_json(suite_path.read_bytes())
                    if suite_path.exists()
                    else None
                )
                for artifact in receipt.artifacts:
                    file = self.output_directory / artifact.filename
                    if Path(artifact.filename).name != artifact.filename or file.is_symlink():
                        continue
                    if not file.is_file() or file.stat().st_size > 20_000_000:
                        continue
                    if hashlib.sha256(file.read_bytes()).hexdigest() != artifact.sha256:
                        continue
                    result.append(
                        {
                            "artifact_id": artifact.sha256[:12],
                            "filename": artifact.filename,
                            "format": artifact.format,
                            "feature": suite.feature_name
                            if suite
                            else re.sub(r"^\d+T\d+Z-", "", stem).replace("-", " "),
                            "cases": artifact.case_count,
                            "created_at": receipt.accepted_at,
                            "accepted_by": receipt.accepted_by,
                            "project": "Unassigned workspace",
                            "storage": "accepted_files",
                            "receipt": path.name,
                        }
                    )
            except (OSError, ValueError):
                continue
        return result

    def entry(self, identifier: str) -> dict | None:
        item = next((r for r in self.inventory() if r["artifact_id"] == identifier), None)
        if not item:
            return None
        stem = item["receipt"].removesuffix("-acceptance.json")
        result = dict(item)
        result["acceptance"] = json.loads((self.output_directory / item["receipt"]).read_text())
        for suffix, key in (("suite", "suite"), ("workflow", "workflow")):
            path = self.output_directory / f"{stem}-{suffix}.json"
            if path.is_file():
                result[key] = json.loads(path.read_text())
        if item["format"] in {"feature", "csv"}:
            result["content"] = (self.output_directory / item["filename"]).read_text()[:100_000]
        if item["format"] == "xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(
                self.output_directory / item["filename"], read_only=True, data_only=True
            )
            try:
                result["rows"] = list(workbook.active.iter_rows(max_row=101, values_only=True))
            finally:
                workbook.close()
        return result

    def _store(
        self, filename: str, format_name: str, count: int, content: bytes
    ) -> AcceptedArtifact:
        destination = self.output_directory / filename
        temporary = destination.with_suffix(destination.suffix + ".tmp")
        temporary.write_bytes(content)
        self._restrict_permissions(temporary, 0o600)
        os.replace(temporary, destination)
        return AcceptedArtifact(
            filename=filename,
            format=format_name,
            case_count=count,
            size_bytes=len(content),
            sha256=hashlib.sha256(content).hexdigest(),
        )

    @staticmethod
    def _stem(feature_name: str, timestamp: datetime) -> str:
        slug = _SAFE_NAME.sub("-", feature_name.casefold()).strip("-")[:60] or "test-suite"
        return f"{timestamp:%Y%m%dT%H%M%S%fZ}-{slug}"

    @staticmethod
    def _restrict_permissions(path: Path, mode: int) -> None:
        try:
            os.chmod(path, mode)
        except OSError:
            pass
